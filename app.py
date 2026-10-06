"""
BC Requirements Analyzer
------------------------
Aplicación Streamlit que analiza requerimientos de Microsoft Dynamics 365
Business Central con Gemini desde tres perspectivas (Consultor Funcional,
Arquitecto AL y QA Tester) y exporta un reporte completo a Word (.docx).

Despliegue: Streamlit Community Cloud (share.streamlit.io)
Secret requerido (opcional): GEMINI_API_KEY
"""

from __future__ import annotations

import os
import re
import time
from datetime import datetime
from io import BytesIO
from typing import Callable, TypeVar

import streamlit as st
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from google import genai
from google.genai import errors as genai_errors
from google.genai import types
from pydantic import BaseModel, Field

# =============================================================================
# Configuración general
# =============================================================================

APP_TITLE = "BC Requirements Analyzer"
DEFAULT_MODEL = "gemini-2.5-flash"
MAX_REQUIREMENT_CHARS = 30_000
MAX_RETRIES = 3
TEMPERATURE = 0.2

st.set_page_config(
    page_title=APP_TITLE,
    page_icon="📊",
    layout="wide",
    initial_sidebar_state="expanded",
)

# =============================================================================
# Esquemas de salida estructurada (Gemini devuelve JSON validado con Pydantic)
# =============================================================================


class AnalisisFuncional(BaseModel):
    objetivo_principal: str = Field(description="Objetivo de negocio del requerimiento en 2-4 oraciones.")
    reglas_negocio: list[str] = Field(description="Reglas de negocio explícitas o derivadas del requerimiento.")
    preguntas_pendientes: list[str] = Field(description="Preguntas que deben resolverse con el usuario clave.")
    informacion_confirmada: list[str] = Field(description="Hechos que el requerimiento afirma explícitamente.")
    supuestos: list[str] = Field(description="Supuestos asumidos por el consultor para poder avanzar.")
    informacion_faltante: list[str] = Field(description="Información necesaria que no fue proporcionada.")


class ObjetoEstandarBC(BaseModel):
    tipo: str = Field(description="Table, Page, Codeunit, Report, Query, XMLport o Enum.")
    id_objeto: str = Field(description="ID estándar del objeto en BC (ej. '36'). Vacío si no hay certeza.")
    nombre: str = Field(description="Nombre estándar exacto en inglés (ej. 'Sales Header').")
    uso: str = Field(description="Cómo interviene en la solución.")


class ObjetoPropuesto(BaseModel):
    tipo: str = Field(description="TableExtension, PageExtension, Codeunit, Table, Page, Report, Enum, EnumExtension.")
    nombre_sugerido: str = Field(description="Nombre sugerido del NUEVO objeto a desarrollar.")
    objeto_base: str = Field(description="Objeto estándar que extiende, o 'N/A' si es objeto nuevo.")
    proposito: str = Field(description="Responsabilidad del objeto.")


class EventoBC(BaseModel):
    objeto_publicador: str = Field(description="Objeto estándar que publica el evento (ej. 'Codeunit 80 Sales-Post').")
    evento: str = Field(description="Nombre del evento de integración o trigger (ej. 'OnBeforePostSalesDoc').")
    proposito: str = Field(description="Lógica a implementar en el subscriber.")


class RiesgoTecnico(BaseModel):
    descripcion: str
    impacto: str = Field(description="Alto, Medio o Bajo.")
    mitigacion: str


class AnalisisTecnico(BaseModel):
    areas_funcionales: list[str] = Field(description="Módulos de BC involucrados (Ventas, Compras, Inventario, Finanzas, etc.).")
    objetos_estandar: list[ObjetoEstandarBC] = Field(description="Solo objetos que EXISTEN en BC estándar.")
    objetos_propuestos: list[ObjetoPropuesto] = Field(description="Extensiones y objetos NUEVOS por desarrollar.")
    eventos_subscribers: list[EventoBC]
    riesgos: list[RiesgoTecnico]
    pasos_desarrollo: list[str] = Field(description="Secuencia ordenada de pasos de desarrollo y despliegue.")
    notas_verificacion: list[str] = Field(
        description="Objetos o eventos cuyo nombre/firma debe validarse contra la versión de BC del cliente."
    )


class CasoPrueba(BaseModel):
    id_caso: str = Field(description="Identificador, ej. 'TC-001'.")
    tipo: str = Field(description="Positivo, Negativo, Límite o Regresión.")
    descripcion: str
    datos_prueba: str
    pasos: str = Field(description="Pasos numerados en una sola cadena.")
    resultado_esperado: str
    prioridad: str = Field(description="Alta, Media o Baja.")


class PlanPruebas(BaseModel):
    escenario: str = Field(description="Escenario principal de prueba end-to-end.")
    datos_necesarios: list[str]
    pasos: list[str] = Field(description="Pasos del escenario principal.")
    resultado_esperado: str
    prioridad: str = Field(description="Alta, Media o Baja.")
    casos: list[CasoPrueba] = Field(
        description="Al menos 2 positivos, 2 negativos, 2 de límite y 2 de regresión."
    )


class Conclusion(BaseModel):
    resumen_ejecutivo: str = Field(description="Párrafo de cierre para un líder de proyecto.")
    complejidad: str = Field(description="Baja, Media o Alta.")
    recomendaciones: list[str]
    siguientes_pasos: list[str]


# =============================================================================
# Prompts de cada rol
# =============================================================================

BASE_RULES = """
Reglas generales:
- Responde SIEMPRE en español profesional.
- Basa tu análisis exclusivamente en el requerimiento y el contexto entregado.
- No inventes datos del cliente; si algo no está, regístralo como supuesto o información faltante.
- Sé concreto y accionable; evita relleno.
""".strip()

PROMPT_FUNCIONAL = f"""
Eres un Consultor Funcional Senior de Microsoft Dynamics 365 Business Central con
más de 10 años de experiencia en implementaciones.

Analiza el requerimiento y entrega: objetivo principal, reglas de negocio,
preguntas pendientes para el usuario clave y una matriz de clasificación que
separe Información confirmada, Supuestos e Información faltante.

{BASE_RULES}
""".strip()

PROMPT_TECNICO = f"""
Eres un Arquitecto de Soluciones AL para Microsoft Dynamics 365 Business Central
(SaaS, extensiones AL, AppSource-ready).

Con base en el requerimiento y el análisis funcional, entrega la propuesta técnica.

Reglas estrictas sobre objetos:
- En 'objetos_estandar' incluye ÚNICAMENTE objetos que existen en Business Central
  estándar (ej. Table 18 Customer, Table 36 Sales Header, Table 37 Sales Line,
  Codeunit 80 Sales-Post, Page 42 Sales Order). Si no estás seguro del ID, deja
  'id_objeto' vacío en lugar de inventarlo.
- NUNCA presentes un objeto personalizado como si fuera estándar. Todo objeto nuevo,
  TableExtension o PageExtension va en 'objetos_propuestos' como desarrollo a crear.
- La personalización debe hacerse vía extensiones y suscripción a eventos; no
  modificar objetos base.
- Prefiere eventos de integración publicados por Microsoft. Si el nombre exacto de un
  evento puede variar entre versiones, agrégalo en 'notas_verificacion'.

{BASE_RULES}
""".strip()

PROMPT_QA = f"""
Eres un QA Tester Senior especializado en ERP (Business Central).

Con base en el requerimiento y los análisis funcional y técnico, diseña el plan de
pruebas: escenario principal, datos necesarios, pasos, resultado esperado y
prioridad, más una tabla completa de casos de prueba que incluya como mínimo
2 casos Positivos, 2 Negativos, 2 de Límite y 2 de Regresión (procesos estándar de
BC que podrían verse afectados por la extensión).

{BASE_RULES}
""".strip()

PROMPT_CONCLUSION = f"""
Eres el Líder de Proyecto de una implementación de Business Central.
Con base en los tres análisis, redacta la conclusión ejecutiva del reporte:
resumen, complejidad estimada, recomendaciones y siguientes pasos.

{BASE_RULES}
""".strip()


# =============================================================================
# Cliente de Gemini
# =============================================================================

T = TypeVar("T", bound=BaseModel)


def get_secret_api_key() -> str | None:
    """Obtiene GEMINI_API_KEY de st.secrets o variable de entorno, sin romper si no existe."""
    try:
        key = st.secrets.get("GEMINI_API_KEY")  # type: ignore[attr-defined]
        if key:
            return str(key).strip()
    except Exception:  # No hay secrets.toml configurado
        pass
    env_key = os.getenv("GEMINI_API_KEY")
    return env_key.strip() if env_key else None


@st.cache_resource(show_spinner=False)
def get_client(api_key: str) -> genai.Client:
    return genai.Client(api_key=api_key)


def call_gemini(
    client: genai.Client,
    model: str,
    system_prompt: str,
    user_content: str,
    schema: type[T],
) -> T:
    """Llama a Gemini con salida JSON estructurada y reintentos con backoff."""
    config = types.GenerateContentConfig(
        system_instruction=system_prompt,
        temperature=TEMPERATURE,
        response_mime_type="application/json",
        response_schema=schema,
    )
    last_error: Exception | None = None
    for attempt in range(1, MAX_RETRIES + 1):
        try:
            response = client.models.generate_content(
                model=model, contents=user_content, config=config
            )
            if isinstance(response.parsed, schema):
                return response.parsed
            if response.text:
                return schema.model_validate_json(response.text)
            raise ValueError("El modelo devolvió una respuesta vacía.")
        except genai_errors.ClientError as exc:
            # 4xx: errores de API key, permisos o request inválido -> no reintentar, salvo 429
            if getattr(exc, "code", None) != 429:
                raise
            last_error = exc
        except (genai_errors.ServerError, ValueError) as exc:
            last_error = exc
        time.sleep(2 ** attempt)
    raise RuntimeError(f"No se obtuvo respuesta válida tras {MAX_RETRIES} intentos: {last_error}")


def friendly_error(exc: Exception) -> str:
    if isinstance(exc, genai_errors.ClientError):
        code = getattr(exc, "code", None)
        if code in (400, 401, 403):
            return "La API Key es inválida o no tiene permisos para este modelo. Verifícala en Google AI Studio."
        if code == 404:
            return "El modelo indicado no existe o no está disponible para tu API Key. Revisa el nombre en el panel lateral."
        if code == 429:
            return "Se alcanzó el límite de cuota de la API de Gemini. Espera un momento e inténtalo de nuevo."
    return f"Ocurrió un error durante el análisis: {exc}"


# =============================================================================
# Orquestación de los análisis
# =============================================================================


def run_analysis(
    client: genai.Client,
    model: str,
    requirement: str,
    progress: Callable[[float, str], None],
) -> dict:
    req_block = f"REQUERIMIENTO ORIGINAL:\n\"\"\"\n{requirement}\n\"\"\""

    progress(0.05, "🧩 Análisis 1/3 — Consultor Funcional BC…")
    funcional = call_gemini(client, model, PROMPT_FUNCIONAL, req_block, AnalisisFuncional)

    progress(0.35, "🏗️ Análisis 2/3 — Arquitecto AL BC…")
    ctx_tecnico = f"{req_block}\n\nANÁLISIS FUNCIONAL (JSON):\n{funcional.model_dump_json(indent=2)}"
    tecnico = call_gemini(client, model, PROMPT_TECNICO, ctx_tecnico, AnalisisTecnico)

    progress(0.65, "🧪 Análisis 3/3 — QA Tester ERP…")
    ctx_qa = f"{ctx_tecnico}\n\nPROPUESTA TÉCNICA (JSON):\n{tecnico.model_dump_json(indent=2)}"
    qa = call_gemini(client, model, PROMPT_QA, ctx_qa, PlanPruebas)

    progress(0.90, "📝 Redactando conclusión…")
    ctx_conc = f"{ctx_qa}\n\nPLAN DE PRUEBAS (JSON):\n{qa.model_dump_json(indent=2)}"
    conclusion = call_gemini(client, model, PROMPT_CONCLUSION, ctx_conc, Conclusion)

    progress(1.0, "✅ Análisis completado")
    return {
        "funcional": funcional,
        "tecnico": tecnico,
        "qa": qa,
        "conclusion": conclusion,
        "model": model,
        "timestamp": datetime.now(),
    }


# =============================================================================
# Generación del documento Word
# =============================================================================

BRAND_COLOR = RGBColor(0x1F, 0x3A, 0x5F)
HEADER_FILL = "1F3A5F"
ZEBRA_FILL = "EEF2F7"


def _shade_cell(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def _add_table(doc: Document, headers: list[str], rows: list[list[str]], widths_cm: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr = table.rows[0]
    _repeat_header(hdr)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _shade_cell(cell, HEADER_FILL)

    for r_idx, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(value or "—"))
            run.font.size = Pt(9)
            if r_idx % 2 == 1:
                _shade_cell(cells[i], ZEBRA_FILL)

    if widths_cm:
        for row in table.rows:
            for i, w in enumerate(widths_cm):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return table


def _add_bullets(doc: Document, items: list[str], empty_text: str = "Sin elementos identificados.") -> None:
    if not items:
        doc.add_paragraph().add_run(empty_text).italic = True
        return
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def _add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Number")


def _set_landscape(section) -> None:
    section.orientation = WD_ORIENT.LANDSCAPE
    w, h = section.page_width, section.page_height
    if w < h:
        section.page_width, section.page_height = h, w


def _set_portrait(section) -> None:
    section.orientation = WD_ORIENT.PORTRAIT
    w, h = section.page_width, section.page_height
    if w > h:
        section.page_width, section.page_height = h, w


def build_docx(project_title: str, requirement: str, results: dict) -> BytesIO:
    funcional: AnalisisFuncional = results["funcional"]
    tecnico: AnalisisTecnico = results["tecnico"]
    qa: PlanPruebas = results["qa"]
    conclusion: Conclusion = results["conclusion"]

    doc = Document()

    # Estilos base
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    for level in (1, 2):
        doc.styles[f"Heading {level}"].font.color.rgb = BRAND_COLOR

    section = doc.sections[0]
    for attr in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, attr, Cm(2.2))

    footer_p = section.footer.paragraphs[0]
    footer_p.text = f"{project_title} — Análisis de requerimiento Business Central"
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_p.runs[0].font.size = Pt(8)

    # Portada / encabezado
    title = doc.add_heading(project_title, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    meta = doc.add_paragraph()
    meta_run = meta.add_run(
        f"Análisis de Requerimiento — Microsoft Dynamics 365 Business Central\n"
        f"Fecha: {results['timestamp']:%d/%m/%Y %H:%M}   |   Modelo: {results['model']}   |   "
        f"Complejidad estimada: {conclusion.complejidad}"
    )
    meta_run.font.size = Pt(9)
    meta_run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    # 1. Requerimiento original
    doc.add_heading("1. Requerimiento original", level=1)
    for para in requirement.strip().split("\n"):
        if para.strip():
            doc.add_paragraph(para.strip())

    # 2. Análisis funcional
    doc.add_heading("2. Análisis funcional", level=1)
    doc.add_heading("2.1 Objetivo principal", level=2)
    doc.add_paragraph(funcional.objetivo_principal)
    doc.add_heading("2.2 Reglas de negocio", level=2)
    _add_bullets(doc, funcional.reglas_negocio)
    doc.add_heading("2.3 Preguntas pendientes", level=2)
    _add_bullets(doc, funcional.preguntas_pendientes, "Sin preguntas pendientes.")
    doc.add_heading("2.4 Matriz de clasificación de información", level=2)
    matriz = (
        [["Información confirmada", x] for x in funcional.informacion_confirmada]
        + [["Supuesto", x] for x in funcional.supuestos]
        + [["Información faltante", x] for x in funcional.informacion_faltante]
    )
    _add_table(doc, ["Clasificación", "Detalle"], matriz, [4.5, 12.0])

    # 3. Propuesta técnica
    doc.add_heading("3. Propuesta técnica", level=1)
    doc.add_heading("3.1 Áreas funcionales involucradas", level=2)
    _add_bullets(doc, tecnico.areas_funcionales)
    doc.add_heading("3.2 Objetos estándar de Business Central", level=2)
    _add_table(
        doc,
        ["Tipo", "ID", "Nombre", "Uso en la solución"],
        [[o.tipo, o.id_objeto, o.nombre, o.uso] for o in tecnico.objetos_estandar],
        [2.5, 1.5, 4.5, 8.0],
    )
    doc.add_heading("3.3 Objetos por desarrollar (extensión AL)", level=2)
    _add_table(
        doc,
        ["Tipo", "Nombre sugerido", "Objeto base", "Propósito"],
        [[o.tipo, o.nombre_sugerido, o.objeto_base, o.proposito] for o in tecnico.objetos_propuestos],
        [3.0, 4.0, 3.5, 6.0],
    )
    doc.add_heading("3.4 Eventos y subscribers", level=2)
    _add_table(
        doc,
        ["Publicador", "Evento", "Lógica del subscriber"],
        [[e.objeto_publicador, e.evento, e.proposito] for e in tecnico.eventos_subscribers],
        [4.5, 4.5, 7.5],
    )
    doc.add_heading("3.5 Pasos de desarrollo", level=2)
    _add_numbered(doc, tecnico.pasos_desarrollo)
    if tecnico.notas_verificacion:
        doc.add_heading("3.6 Puntos a verificar contra la versión de BC del cliente", level=2)
        _add_bullets(doc, tecnico.notas_verificacion)

    # 4. Riesgos
    doc.add_heading("4. Riesgos", level=1)
    _add_table(
        doc,
        ["Riesgo", "Impacto", "Mitigación"],
        [[r.descripcion, r.impacto, r.mitigacion] for r in tecnico.riesgos],
        [7.0, 2.0, 7.5],
    )

    # 5. Plan de pruebas (sección horizontal para la tabla de casos)
    test_section = doc.add_section()
    _set_landscape(test_section)
    doc.add_heading("5. Plan de pruebas", level=1)
    doc.add_heading("5.1 Escenario principal", level=2)
    _add_table(
        doc,
        ["Campo", "Detalle"],
        [
            ["Escenario", qa.escenario],
            ["Datos necesarios", "\n".join(f"• {d}" for d in qa.datos_necesarios)],
            ["Pasos", "\n".join(f"{i}. {p}" for i, p in enumerate(qa.pasos, 1))],
            ["Resultado esperado", qa.resultado_esperado],
            ["Prioridad", qa.prioridad],
        ],
        [4.5, 19.0],
    )
    doc.add_heading("5.2 Casos de prueba", level=2)
    _add_table(
        doc,
        ["ID", "Tipo", "Descripción", "Datos de prueba", "Pasos", "Resultado esperado", "Prioridad"],
        [
            [c.id_caso, c.tipo, c.descripcion, c.datos_prueba, c.pasos, c.resultado_esperado, c.prioridad]
            for c in qa.casos
        ],
        [1.6, 2.0, 4.2, 3.8, 5.5, 4.5, 1.9],
    )

    # 6. Conclusión
    concl_section = doc.add_section()
    _set_portrait(concl_section)
    doc.add_heading("6. Conclusión", level=1)
    doc.add_paragraph(conclusion.resumen_ejecutivo)
    doc.add_heading("6.1 Recomendaciones", level=2)
    _add_bullets(doc, conclusion.recomendaciones)
    doc.add_heading("6.2 Siguientes pasos", level=2)
    _add_numbered(doc, conclusion.siguientes_pasos)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def safe_filename(text: str) -> str:
    slug = re.sub(r"[^\w\-]+", "_", text.strip(), flags=re.UNICODE).strip("_")
    return (slug or "Analisis_BC")[:60]


# =============================================================================
# Renderizado de resultados en pantalla
# =============================================================================


def _bullets_md(items: list[str], empty: str = "_Sin elementos identificados._") -> str:
    return "\n".join(f"- {x}" for x in items) if items else empty


def render_funcional(f: AnalisisFuncional) -> None:
    st.subheader("🎯 Objetivo principal")
    st.write(f.objetivo_principal)
    c1, c2 = st.columns(2)
    with c1:
        st.subheader("📏 Reglas de negocio")
        st.markdown(_bullets_md(f.reglas_negocio))
    with c2:
        st.subheader("❓ Preguntas pendientes")
        st.markdown(_bullets_md(f.preguntas_pendientes, "_Sin preguntas pendientes._"))
    st.subheader("🗂️ Matriz de clasificación")
    m1, m2, m3 = st.columns(3)
    with m1:
        st.success("**Información confirmada**")
        st.markdown(_bullets_md(f.informacion_confirmada))
    with m2:
        st.warning("**Supuestos**")
        st.markdown(_bullets_md(f.supuestos))
    with m3:
        st.error("**Información faltante**")
        st.markdown(_bullets_md(f.informacion_faltante))


def render_tecnico(t: AnalisisTecnico) -> None:
    st.subheader("🧭 Áreas funcionales")
    st.markdown(" · ".join(f"`{a}`" for a in t.areas_funcionales) or "_N/A_")

    st.subheader("📦 Objetos estándar de BC")
    st.dataframe(
        [{"Tipo": o.tipo, "ID": o.id_objeto, "Nombre": o.nombre, "Uso": o.uso} for o in t.objetos_estandar],
        width="stretch", hide_index=True,
    )
    st.subheader("🛠️ Objetos por desarrollar (extensión AL)")
    st.caption("Objetos nuevos a crear — no existen en BC estándar.")
    st.dataframe(
        [
            {"Tipo": o.tipo, "Nombre sugerido": o.nombre_sugerido, "Objeto base": o.objeto_base, "Propósito": o.proposito}
            for o in t.objetos_propuestos
        ],
        width="stretch", hide_index=True,
    )
    st.subheader("⚡ Eventos / Subscribers")
    st.dataframe(
        [{"Publicador": e.objeto_publicador, "Evento": e.evento, "Lógica": e.proposito} for e in t.eventos_subscribers],
        width="stretch", hide_index=True,
    )
    st.subheader("⚠️ Riesgos")
    st.dataframe(
        [{"Riesgo": r.descripcion, "Impacto": r.impacto, "Mitigación": r.mitigacion} for r in t.riesgos],
        width="stretch", hide_index=True,
    )
    st.subheader("🪜 Pasos de desarrollo")
    st.markdown("\n".join(f"{i}. {p}" for i, p in enumerate(t.pasos_desarrollo, 1)))
    if t.notas_verificacion:
        with st.expander("🔍 Puntos a verificar contra la versión de BC del cliente"):
            st.markdown(_bullets_md(t.notas_verificacion))


def render_qa(q: PlanPruebas) -> None:
    c1, c2 = st.columns([3, 1])
    with c1:
        st.subheader("🎬 Escenario principal")
        st.write(q.escenario)
    with c2:
        st.metric("Prioridad", q.prioridad)
        st.metric("Casos de prueba", len(q.casos))
    c3, c4 = st.columns(2)
    with c3:
        st.subheader("🧾 Datos necesarios")
        st.markdown(_bullets_md(q.datos_necesarios))
    with c4:
        st.subheader("👣 Pasos")
        st.markdown("\n".join(f"{i}. {p}" for i, p in enumerate(q.pasos, 1)))
    st.subheader("✅ Resultado esperado")
    st.info(q.resultado_esperado)

    st.subheader("📋 Tabla de casos de prueba")
    tipos = sorted({c.tipo for c in q.casos})
    seleccion = st.multiselect("Filtrar por tipo", tipos, default=tipos, key="qa_filter")
    st.dataframe(
        [
            {
                "ID": c.id_caso, "Tipo": c.tipo, "Descripción": c.descripcion, "Datos": c.datos_prueba,
                "Pasos": c.pasos, "Resultado esperado": c.resultado_esperado, "Prioridad": c.prioridad,
            }
            for c in q.casos if c.tipo in seleccion
        ],
        width="stretch", hide_index=True,
    )


# =============================================================================
# Interfaz principal
# =============================================================================


def sidebar() -> tuple[str | None, str, str]:
    with st.sidebar:
        st.header("⚙️ Configuración")

        # La API Key se toma del Secret GEMINI_API_KEY; el usuario final nunca la ve ni la captura.
        api_key = get_secret_api_key()
        if api_key:
            st.success("✅ Conectado a Gemini")
        else:
            st.error(
                "Falta configurar el Secret GEMINI_API_KEY en Streamlit Cloud "
                "(Settings → Secrets). Contacta al administrador."
            )

        st.divider()
        project_title = st.text_input("Título del proyecto", value="Análisis de Requerimiento BC")
        with st.expander("Opciones avanzadas"):
            model = st.text_input(
                "Modelo de Gemini", value=DEFAULT_MODEL,
                help="Cámbialo si Google retira o renombra el modelo.",
            ).strip() or DEFAULT_MODEL

        st.divider()
        if st.button("🧹 Limpiar resultados", width="stretch"):
            st.session_state.pop("results", None)
            st.rerun()
        st.caption("Los análisis son generados por IA y deben validarse por un consultor antes de usarse con el cliente.")
    return api_key, project_title.strip() or "Análisis de Requerimiento BC", model


def main() -> None:
    api_key, project_title, model = sidebar()

    st.title("📊 Analizador de Requerimientos — Business Central")
    st.caption("Consultor Funcional · Arquitecto AL · QA Tester ERP — con exportación a Word")

    requirement = st.text_area(
        "Requerimiento",
        height=260,
        max_chars=MAX_REQUIREMENT_CHARS,
        placeholder=(
            "Ejemplo: Al registrar un pedido de venta, el sistema debe bloquear la liberación si el "
            "cliente tiene facturas vencidas a más de 30 días, salvo que un usuario con rol de "
            "Crédito autorice la excepción..."
        ),
        key="requirement",
    )

    run = st.button("🚀 Ejecutar Análisis", type="primary", width="stretch")

    if run:
        if not api_key:
            st.error("La app no tiene configurada la API Key de Gemini. Contacta al administrador.")
            st.stop()
        if len(requirement.strip()) < 30:
            st.warning("Describe el requerimiento con más detalle (mínimo 30 caracteres).")
            st.stop()

        bar = st.progress(0.0, text="Iniciando…")
        try:
            results = run_analysis(
                get_client(api_key), model, requirement.strip(),
                lambda pct, msg: bar.progress(pct, text=msg),
            )
            results["requirement"] = requirement.strip()
            st.session_state["results"] = results
        except Exception as exc:  # noqa: BLE001 - se muestra al usuario de forma controlada
            bar.empty()
            st.error(friendly_error(exc))
            st.stop()
        bar.empty()
        st.toast("Análisis completado", icon="✅")

    results = st.session_state.get("results")
    if not results:
        st.info("Pega un requerimiento y presiona **🚀 Ejecutar Análisis** para comenzar.")
        return

    st.divider()
    tab1, tab2, tab3, tab4 = st.tabs(
        ["🧩 Análisis Funcional", "🏗️ Propuesta Técnica AL", "🧪 Plan de Pruebas QA", "📝 Conclusión"]
    )
    with tab1:
        render_funcional(results["funcional"])
    with tab2:
        render_tecnico(results["tecnico"])
    with tab3:
        render_qa(results["qa"])
    with tab4:
        c: Conclusion = results["conclusion"]
        st.metric("Complejidad estimada", c.complejidad)
        st.write(c.resumen_ejecutivo)
        st.subheader("Recomendaciones")
        st.markdown(_bullets_md(c.recomendaciones))
        st.subheader("Siguientes pasos")
        st.markdown("\n".join(f"{i}. {p}" for i, p in enumerate(c.siguientes_pasos, 1)))

    st.divider()
    docx_buffer = build_docx(project_title, results["requirement"], results)
    st.download_button(
        "📥 Descargar reporte completo (.docx)",
        data=docx_buffer,
        file_name=f"{safe_filename(project_title)}_{results['timestamp']:%Y%m%d_%H%M}.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        type="primary",
        width="stretch",
    )


if __name__ == "__main__":
    main()
